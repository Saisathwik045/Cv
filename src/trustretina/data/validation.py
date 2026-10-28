"""Read-only original ODIR workbook/image validation; labels remain patient-derived."""
import hashlib
import json
import math
from pathlib import Path
import warnings

from trustretina.config import CLASSES


class DataValidationError(ValueError):
    """A dataset cannot safely enter a research split."""


def canonical_hash(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, separators=(',', ':'),
                                     allow_nan=False).encode()).hexdigest()


def file_hash(path):
    digest = hashlib.sha256()
    with Path(path).open('rb') as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def inside(root, value):
    path = (root / value).resolve()
    if not path.is_relative_to(root):
        raise DataValidationError('Input path escapes the configured dataset root')
    return path


def discover(config):
    root = Path(config['data']['root']).resolve()
    if not root.is_dir():
        raise DataValidationError('Configured dataset root does not exist')
    workbooks = sorted(str(p.relative_to(root)) for p in root.rglob('data.xlsx')
                       if p.is_file() and p.resolve().is_relative_to(root))
    image_dirs = sorted(str(p.relative_to(root)) for p in root.rglob('*')
                        if p.is_dir() and p.name == Path(config['data']['images']).name
                        and p.resolve().is_relative_to(root))
    return {'dataset_root': str(root), 'authoritative_workbook_candidates': workbooks,
            'image_directory_candidates': image_dirs,
            'automatic_selection': False, 'label_source': 'patient'}


def patient_id(value):
    # Numeric Excel IDs and textual numeric IDs canonicalize identically.
    if isinstance(value, bool) or value is None:
        raise DataValidationError('Missing/invalid patient ID')
    if isinstance(value, (int, float)):
        if not math.isfinite(value) or value < 0 or int(value) != value:
            raise DataValidationError('Patient ID must be a nonnegative integer or nonempty text')
        return str(int(value))
    if not isinstance(value, str) or not value.strip() or value.startswith('='):
        raise DataValidationError('Missing/invalid patient ID')
    value = value.strip()
    return str(int(value)) if value.isdecimal() else value


def parse_workbook(path, sheet_name=None, data_kind='original_odir'):
    from openpyxl import load_workbook
    required = ['ID', 'Left-Fundus', 'Right-Fundus', *CLASSES]
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        if 'TRUSTRETINA_SYNTHETIC_FIXTURE' in book.sheetnames and data_kind != 'synthetic_fixture':
            raise DataValidationError('Synthetic fixture must be explicitly marked in data.kind')
        candidates = []
        for sheet in book:
            header = list(next(sheet.iter_rows(values_only=True), ()))
            if all(name in header for name in required):
                candidates.append(sheet.title)
        if sheet_name is None:
            if len(candidates) != 1:
                raise DataValidationError('Require exactly one label sheet or configure data.sheet')
            sheet_name = candidates[0]
        if sheet_name not in candidates:
            raise DataValidationError('Configured sheet lacks required ODIR columns')
        rows = book[sheet_name].iter_rows(values_only=True)
        header = list(next(rows))
        if any(header.count(name) != 1 for name in required):
            raise DataValidationError('Duplicate required workbook column')
        columns = {name: header.index(name) for name in required}
        patients, seen, references = [], set(), set()
        for number, row in enumerate(rows, 2):
            if not any(value is not None for value in row):
                continue
            try:
                pid = patient_id(row[columns['ID']])
                if pid in seen:
                    raise DataValidationError('Duplicate patient ID')
                labels = [row[columns[name]] for name in CLASSES]
                if any(isinstance(v, bool) or not isinstance(v, (int, float)) or v not in (0, 1)
                       for v in labels):
                    raise DataValidationError('Require eight numeric binary nonmissing labels; formulas rejected')
                labels = [int(v) for v in labels]
                if not any(labels) or (labels[0] and any(labels[1:])):
                    raise DataValidationError('All-zero or contradictory N/disease labels require adjudication')
                eyes = {}
                for eye, field in [('left', 'Left-Fundus'), ('right', 'Right-Fundus')]:
                    name = row[columns[field]]
                    if not isinstance(name, str) or not name.strip() or name.startswith('='):
                        raise DataValidationError('Both eye filenames are required')
                    name = name.strip()
                    if Path(name).stem != f'{pid}_{eye}':
                        raise DataValidationError('Eye filename does not match source patient ID/eye')
                    if name in references:
                        raise DataValidationError('Duplicate image reference')
                    references.add(name)
                    eyes[eye] = name
                patients.append({'patient_id': pid, 'labels': labels, 'eyes': eyes})
                seen.add(pid)
            except DataValidationError as exc:
                # Avoid emitting identifiers or labels into CI/logs.
                raise DataValidationError(f'Workbook row {number}: {exc}') from None
        if not patients:
            raise DataValidationError('No labeled patients in workbook')
        return sorted(patients, key=lambda p: p['patient_id']), sheet_name
    finally:
        book.close()


def image_record(path, root):
    from PIL import Image, ImageOps
    original_sha256 = file_hash(path)
    with warnings.catch_warnings():
        warnings.simplefilter('error', Image.DecompressionBombWarning)
        with Image.open(path) as image:
            image.verify()
        with Image.open(path) as image:
            image = ImageOps.exif_transpose(image).convert('RGB')
            image.load()
            if min(image.size) < 2:
                raise DataValidationError('Image dimensions too small')
            # dHash is a duplicate-screening proxy, not a clinical quality metric.
            small = image.convert('L').resize((9, 8), Image.Resampling.LANCZOS)
            pixels = list(small.tobytes())
            dhash = 0
            for y in range(8):
                for x in range(8):
                    dhash = (dhash << 1) | int(pixels[y * 9 + x] > pixels[y * 9 + x + 1])
            if file_hash(path) != original_sha256:
                raise DataValidationError('Image changed during validation')
            return {'path': str(path.relative_to(root)), 'sha256': original_sha256,
                    'pixel_sha256': hashlib.sha256(str(image.size).encode() + image.tobytes()).hexdigest(),
                    'width': image.width, 'height': image.height, 'dhash': f'{dhash:016x}'}


def validate_dataset(config):
    root = Path(config['data']['root']).resolve()
    workbook = inside(root, config['data']['labels'])
    images = inside(root, config['data']['images'])
    if workbook.name != 'data.xlsx' or not workbook.is_file() or not images.is_dir():
        raise DataValidationError('Configured original data.xlsx or image directory is unavailable; run discover-data')
    before = file_hash(workbook)
    patients, sheet = parse_workbook(workbook, config['data'].get('sheet'), config['data'].get('kind', 'original_odir'))
    threshold = config['data'].get('near_duplicate_hamming', 2)
    if type(threshold) is not int or not 0 <= threshold <= 8:
        raise DataValidationError('near_duplicate_hamming must be an integer from 0 to 8')
    resolved_seen = set()
    for patient in patients:
        records = {}
        for eye, name in patient['eyes'].items():
            path = inside(images, name)
            if not path.is_file():
                raise DataValidationError('A referenced eye image is missing')
            if path in resolved_seen:
                raise DataValidationError('Multiple eye references resolve to the same file')
            resolved_seen.add(path)
            try:
                record = image_record(path, root)
            except (OSError, ValueError, SyntaxError, Warning) as exc:
                raise DataValidationError(f'Image decoding failed ({type(exc).__name__})') from None
            records[eye] = record
        patient['eyes'] = records
    if file_hash(workbook) != before:
        raise DataValidationError('Workbook changed during validation')
    inventory = {'schema_version': 1, 'label_source': 'patient', 'class_order': CLASSES,
                 'data_kind': config['data'].get('kind', 'original_odir'),
                 'workbook': {'path': str(workbook.relative_to(root)), 'sha256': before, 'sheet': sheet},
                 'near_duplicate_hamming': threshold, 'patients': patients}
    from .duplicates import POLICY, adjudicate
    policy = config['data'].get('duplicate_policy', 'reject')
    if policy not in ('reject', POLICY):
        raise DataValidationError('Unknown duplicate handling policy')
    adjudication = adjudicate(patients, threshold, root)
    if policy == 'reject' and adjudication['pairs']:
        exact = adjudication['counts']['confirmed_exact_pairs']
        message = ('Exact file/pixel duplicate requires adjudication before splitting' if exact
                   else 'Potential near-duplicate requires adjudication before splitting')
        raise DataValidationError(message)
    if file_hash(workbook) != before:
        raise DataValidationError('Workbook changed during duplicate comparison')
    inventory['duplicate_policy'] = policy
    inventory['duplicate_adjudication'] = adjudication
    inventory['fingerprint'] = canonical_hash(inventory)
    return inventory
