"""Real XLSX/image I/O on explicitly synthetic, runtime-generated fixtures."""
import copy
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest

from openpyxl import load_workbook
from trustretina.config import PARTITIONS
from trustretina.data.validation import DataValidationError, discover, parse_workbook, validate_dataset, file_hash
from trustretina.data.splits import prepare_data, make_splits, audit_manifest
from synthetic import make_fixture


class DataTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.config = make_fixture(self.temp.name)
        self.root = Path(self.config['data']['root'])
        self.workbook = self.root / 'data.xlsx'
        self.output = Path(self.temp.name) / 'artifacts' / 'bundle'

    def edit(self, cell, value):
        book = load_workbook(self.workbook)
        book['SyntheticLabels'][cell] = value
        book.save(self.workbook)
        book.close()

    def test_cli_discovery_preparation_and_reuse(self):
        source = Path(__file__).resolve().parents[1] / 'configs/phase1.toml'
        text = source.read_text().replace('root = "../data/raw/ODIR-5K"',
                                          'root = ' + json.dumps(str(self.root)))
        text = text.replace('kind = "original_odir"', 'kind = "synthetic_fixture"')
        config_path = Path(self.temp.name) / 'fixture.toml'
        config_path.write_text(text)
        for command in ('discover-data', 'prepare-data', 'prepare-data'):
            args = [sys.executable, '-m', 'trustretina', command, '--config', str(config_path)]
            if command == 'prepare-data':
                args += ['--output', str(self.output)]
            result = subprocess.run(args, capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            report = json.loads(result.stdout)
            if command == 'prepare-data':
                self.assertEqual(report['data_kind'], 'synthetic_fixture')
        self.assertTrue(report['reused'])

    def test_discovery_does_not_select_ambiguous_workbooks(self):
        found = discover(self.config)
        self.assertEqual(found['authoritative_workbook_candidates'], ['data.xlsx'])
        self.assertFalse(found['automatic_selection'])
        (self.root / 'other').mkdir()
        shutil.copy(self.workbook, self.root / 'other' / 'data.xlsx')
        self.assertEqual(len(discover(self.config)['authoritative_workbook_candidates']), 2)

    def test_frozen_manifests_support_determinism_raw_immutability(self):
        originals = {p: file_hash(p) for p in self.root.rglob('*') if p.is_file()}
        first = prepare_data(self.config, self.output)
        self.assertFalse(first['reused'])
        self.assertEqual(first['data_kind'], 'synthetic_fixture')
        for role in PARTITIONS:
            self.assertEqual(set(first['support'][role]['positive_patients']), set('NDGCAHMO'))
            self.assertTrue(all(v > 0 for v in first['support'][role]['positive_patients'].values()))
            self.assertTrue(all(v > 0 for v in first['support'][role]['negative_patients'].values()))
        hashes = {p.name: file_hash(p) for p in self.output.iterdir()}
        self.assertTrue(prepare_data(self.config, self.output)['reused'])
        self.assertEqual(hashes, {p.name: file_hash(p) for p in self.output.iterdir()})
        second = self.output.with_name('independent')
        prepare_data(self.config, second)
        self.assertEqual(hashes, {p.name: file_hash(p) for p in second.iterdir()})
        self.assertEqual(originals, {p: file_hash(p) for p in originals})
        inventory = json.loads((self.output / 'inventory.json').read_text())
        manifest = json.loads((self.output / 'manifest.json').read_text())
        self.assertEqual(audit_manifest(manifest, inventory), manifest['support'])
        ids = [p['patient_id'] for rows in manifest['partitions'].values() for p in rows]
        self.assertEqual(len(ids), len(set(ids)))
        self.assertTrue(all(set(p['eyes']) == {'left', 'right'} for rows in manifest['partitions'].values() for p in rows))

    def test_multilabel_patient_targets_preserved(self):
        # Patient 1 is D-positive; add G without changing source granularity.
        self.edit('F3', 1)
        rows, _ = parse_workbook(self.workbook, data_kind='synthetic_fixture')
        self.assertEqual(next(p['labels'] for p in rows if p['patient_id'] == '1'), [0, 1, 1, 0, 0, 0, 0, 0])
        self.assertFalse(prepare_data(self.config, self.output)['reused'])

    def test_fail_malformed_labels_ids_and_eye_association(self):
        for cell, value in [('D2', None), ('D2', '=1'), ('D2', 2), ('D2', '1'),
                            ('E2', 1), ('A3', 0), ('B2', '1_left.png'),
                            ('C2', '0_left.png'), ('A2', None)]:
            with self.subTest(cell=cell, value=value):
                original = self.workbook.read_bytes()
                self.edit(cell, value)
                with self.assertRaises(DataValidationError):
                    validate_dataset(self.config)
                self.workbook.write_bytes(original)

    def test_missing_and_corrupt_images(self):
        path = self.root / 'Training Images/0_left.png'
        path.unlink()
        with self.assertRaises(DataValidationError):
            prepare_data(self.config, self.output)
        self.assertFalse(self.output.exists())
        path.write_text('not an image')
        with self.assertRaises(DataValidationError):
            validate_dataset(self.config)

    def test_exact_and_near_duplicates_block_split(self):
        images = self.root / 'Training Images'
        shutil.copy(images / '0_left.png', images / '1_left.png')
        with self.assertRaisesRegex(DataValidationError, 'duplicate'):
            validate_dataset(self.config)
        from PIL import Image
        # Re-encode identical pixels as another file format, preserving the filename.
        with Image.open(images / '0_left.png') as image:
            image.save(images / '1_left.png', format='BMP')
        with self.assertRaisesRegex(DataValidationError, 'duplicate'):
            validate_dataset(self.config)

    def test_perceptual_near_duplicate_requires_adjudication(self):
        from PIL import Image
        images = self.root / 'Training Images'
        with Image.open(images / '0_left.png') as original:
            altered = original.copy()
            pixel = altered.getpixel((0, 0))
            altered.putpixel((0, 0), tuple((v + 1) % 256 for v in pixel))
            altered.save(images / '1_left.png')
        with self.assertRaisesRegex(DataValidationError, 'near-duplicate'):
            validate_dataset(self.config)

    def test_unsafe_paths_and_symlinks(self):
        self.config['data']['labels'] = '../data.xlsx'
        with self.assertRaises(DataValidationError):
            validate_dataset(self.config)
        self.config['data']['labels'] = 'data.xlsx'
        path = self.root / 'Training Images/0_left.png'
        path.unlink()
        outside = Path(self.temp.name) / 'outside.png'
        outside.write_bytes(b'not used')
        path.symlink_to(outside)
        with self.assertRaises(DataValidationError):
            validate_dataset(self.config)

    def test_synthetic_marker_cannot_be_reported_as_real(self):
        self.config['data']['kind'] = 'original_odir'
        with self.assertRaisesRegex(DataValidationError, 'Synthetic'):
            validate_dataset(self.config)

    def test_existing_bundle_changed_inputs_or_corruption_rejected(self):
        prepare_data(self.config, self.output)
        original = (self.output / 'train.json').read_bytes()
        (self.output / 'train.json').write_text('[]')
        with self.assertRaisesRegex(DataValidationError, 'checksum'):
            prepare_data(self.config, self.output)
        (self.output / 'train.json').write_bytes(original)
        self.config['seed'] += 1
        with self.assertRaisesRegex(DataValidationError, 'different inputs'):
            prepare_data(self.config, self.output)
        self.assertEqual((self.output / 'train.json').read_bytes(), original)

    def test_global_and_partition_support_fail_closed(self):
        inventory = validate_dataset(self.config)
        for p in inventory['patients']:
            p['labels'][7] = 0
        with self.assertRaisesRegex(DataValidationError, 'global'):
            make_splits(inventory, self.config)
        inventory = validate_dataset(self.config)
        config = copy.deepcopy(self.config)
        config['split']['minimum_positive_patients'] = 3
        repaired = make_splits(inventory, config)
        self.assertTrue(all(min(role['positive_patients'].values()) >= 3 for role in repaired['support'].values()))
        config['split']['minimum_positive_patients'] = 6
        with self.assertRaises(DataValidationError):
            make_splits(inventory, config)

    def test_tampered_patient_group_or_labels_rejected(self):
        inventory = validate_dataset(self.config)
        original = make_splits(inventory, self.config)
        for mutation in ('leak', 'missing_eye', 'label'):
            manifest = copy.deepcopy(original)
            row = manifest['partitions']['train'][0]
            if mutation == 'leak':
                manifest['partitions']['validation'].append(row)
            elif mutation == 'missing_eye':
                del row['eyes']['right']
            else:
                row['labels'] = [1] * 8
            with self.assertRaises(DataValidationError):
                audit_manifest(manifest, inventory)

    def test_multiple_label_sheets_require_explicit_selection(self):
        book = load_workbook(self.workbook)
        book.copy_worksheet(book['SyntheticLabels'])
        book.save(self.workbook)
        book.close()
        with self.assertRaisesRegex(DataValidationError, 'exactly one'):
            validate_dataset(self.config)
        self.config['data']['sheet'] = 'SyntheticLabels'
        self.assertTrue(validate_dataset(self.config)['patients'])
